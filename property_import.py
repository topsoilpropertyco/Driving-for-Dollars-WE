"""Synthetic-only property import planning and validation.

This module intentionally has no database, network, or file-write behavior.
It converts a CSV export into an import *plan* so that later protected storage
can apply reviewed records atomically.  Diagnostics identify row numbers and
field names only; they never echo a property, owner, or contact value.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook


HEADER_ALIASES = {
    "apn": {"apn", "apn id", "parcel", "parcel id", "parcel number", "tax id"},
    "address": {"address", "property address", "site address", "street address"},
    "unit": {"unit", "unit #", "apt", "apartment", "suite"},
    "city": {"city", "property city"},
    "state": {"state", "property state"},
    "zip": {"zip", "zip code", "postal code"},
    "county": {"county", "property county"},
}


def _text(value: object | None) -> str:
    return " ".join(str(value or "").strip().split())


def _key(value: object | None) -> str:
    return re.sub(r"[^A-Z0-9]", "", _text(value).upper())


def _header(value: object | None) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", _text(value).lower()).split())


def resolve_headers(fieldnames: Iterable[str | None]) -> dict[str, str]:
    """Map recognized canonical fields to the export's exact headings."""
    normalized = {_header(name): _text(name) for name in fieldnames if _text(name)}
    result: dict[str, str] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                result[canonical] = normalized[alias]
                break
    return result


@dataclass(frozen=True)
class PlannedRecord:
    row_number: int
    identity_kind: str  # apn | address_candidate
    identity_key: str
    review_required: bool
    source_fields_present: tuple[str, ...]


@dataclass(frozen=True)
class RejectedRow:
    row_number: int
    code: str
    fields: tuple[str, ...]


@dataclass(frozen=True)
class ImportPlan:
    source_name: str
    source_fingerprint: str
    recognized_headers: tuple[str, ...]
    records: tuple[PlannedRecord, ...]
    rejected: tuple[RejectedRow, ...]

    def summary(self) -> dict[str, object]:
        reasons = Counter(row.code for row in self.rejected)
        return {
            "source_name": self.source_name,
            "source_fingerprint": self.source_fingerprint,
            "recognized_headers": list(self.recognized_headers),
            "accepted": len(self.records),
            "rejected": len(self.rejected),
            "rejections_by_code": dict(sorted(reasons.items())),
            "review_required": sum(record.review_required for record in self.records),
        }


def _fingerprint(
    source_name: str,
    fieldnames: Iterable[str | None],
    records: Iterable[PlannedRecord],
    rejected: Iterable[RejectedRow],
) -> str:
    """Create an idempotency key from normalized identities, never raw provider fields."""
    planned = sorted(
        f"accept:{record.row_number}:{record.identity_kind}:{record.identity_key}:{int(record.review_required)}"
        for record in records
    )
    refused = sorted(
        f"reject:{row.row_number}:{row.code}:{','.join(row.fields)}" for row in rejected
    )
    material = "\x1f".join([source_name, *sorted(_header(name) for name in fieldnames), *planned, *refused])
    return hashlib.sha256(material.encode("utf-8")).hexdigest()[:16]


def plan_rows(source_name: str, fieldnames: Iterable[str | None], rows: Iterable[Mapping[str, object | None]]) -> ImportPlan:
    """Validate identity candidates without persisting or disclosing row values."""
    names = list(fieldnames)
    headers = resolve_headers(names)
    accepted: list[PlannedRecord] = []
    rejected: list[RejectedRow] = []
    seen: set[tuple[str, str]] = set()
    row_count = 0

    for row_number, row in enumerate(rows, start=2):
        row_count += 1
        apn = _key(row.get(headers.get("apn", "")))
        county = _key(row.get(headers.get("county", "")))
        state = _key(row.get(headers.get("state", "")))
        address = _key(row.get(headers.get("address", "")))
        unit = _key(row.get(headers.get("unit", "")))
        city = _key(row.get(headers.get("city", "")))
        if apn and county and state:
            kind, identity, review = "apn", f"{state}:{county}:{apn}", False
        elif address and city and state:
            # Never silently merge address-only records into a property.
            kind, identity, review = "address_candidate", f"{state}:{city}:{address}:{unit}", True
        else:
            rejected.append(RejectedRow(row_number, "missing_stable_identity", ("APN", "county", "state", "address", "city")))
            continue
        duplicate_key = (kind, identity)
        if duplicate_key in seen:
            rejected.append(RejectedRow(row_number, "duplicate_identity_in_file", ("APN", "address")))
            continue
        seen.add(duplicate_key)
        present = tuple(sorted(_header(name) for name in names if _text(row.get(name))))
        accepted.append(PlannedRecord(row_number, kind, identity, review, present))

    return ImportPlan(
        source_name=source_name,
        source_fingerprint=_fingerprint(source_name, names, accepted, rejected),
        recognized_headers=tuple(sorted(headers)),
        records=tuple(accepted),
        rejected=tuple(rejected),
    )


def plan_csv(path: Path, source_name: str) -> ImportPlan:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError("CSV has no header row")
        return plan_rows(source_name, reader.fieldnames, reader)


def plan_xlsx(path: Path, source_name: str, sheet_name: str | None = None) -> ImportPlan:
    """Plan one XLSX worksheet in read-only mode without writing a workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or not any(_text(value) for value in header):
            raise ValueError("XLSX sheet has no header row")
        names = [_text(value) for value in header]
        mapped_rows = ({name: value for name, value in zip(names, values)} for values in rows)
        return plan_rows(source_name, names, mapped_rows)
    finally:
        workbook.close()


def plan_file(path: Path, source_name: str, sheet_name: str | None = None) -> ImportPlan:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return plan_csv(path, source_name)
    if suffix == ".xlsx":
        return plan_xlsx(path, source_name, sheet_name)
    raise ValueError("only CSV and XLSX files are supported")


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate a property CSV/XLSX file without importing it")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--source", required=True, help="Provider label, e.g. propstream-export")
    parser.add_argument("--sheet", help="XLSX worksheet name; defaults to the active worksheet")
    args = parser.parse_args()
    plan = plan_file(args.input, args.source, args.sheet)
    # Deliberately output aggregate diagnostics only.
    print(json.dumps(plan.summary(), indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
